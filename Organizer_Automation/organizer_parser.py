import openpyxl
import os
import re

def clean_val(v):
    if v is None:
        return ""
    s = str(v).strip()
    return s

def clean_ssn(v):
    if not v:
        return ""
    digits = re.sub(r"[^0-9]", "", str(v))
    if len(digits) == 9:
        return f"{digits[:3]}-{digits[3:5]}-{digits[5:]}"
    return str(v).strip()

def clean_phone(v):
    if not v:
        return ""
    digits = re.sub(r"[^0-9]", "", str(v))
    if len(digits) == 10:
        return f"{digits[:3]}-{digits[3:6]}-{digits[6:]}"
    return str(v).strip()

def parse_organizer(file_path):
    """
    Parses the 2025 Personal Tax Organizer Excel workbook into a structured dictionary.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    wb = openpyxl.load_workbook(file_path, data_only=True)
    if "Organizer" not in wb.sheetnames:
        raise ValueError("Workbook must contain an 'Organizer' sheet.")

    ws = wb["Organizer"]

    data = {
        "taxpayer": {},
        "spouse": {},
        "address": {},
        "banking": {},
        "dependents": [],
        "residency_questions": {},
        "income": {},
        "adjustments": {},
        "deductions": {}
    }

    # Helper to get cell value safely
    def cell(r, c):
        return clean_val(ws.cell(r, c).value)

    # 1. Taxpayer Information
    # Row 18: Col 5=First, Col 9=Middle, Col 13=Last, Col 17=SSN
    data["taxpayer"]["first_name"] = cell(18, 5)
    data["taxpayer"]["middle_name"] = cell(18, 9)
    data["taxpayer"]["last_name"] = cell(18, 13)
    data["taxpayer"]["ssn"] = clean_ssn(cell(18, 17))
    # Row 21: Col 3=Occupation, Col 8=DOB
    data["taxpayer"]["occupation"] = cell(21, 3)
    data["taxpayer"]["dob"] = cell(21, 8)
    # Row 22: Col 3=Main Phone, Col 9=Other Phone, Col 15=Email
    data["taxpayer"]["main_phone"] = clean_phone(cell(22, 3))
    data["taxpayer"]["other_phone"] = clean_phone(cell(22, 9))
    data["taxpayer"]["email"] = cell(22, 15)

    # 2. Spouse Information
    # Row 19: Col 5=First, Col 9=Middle, Col 13=Last, Col 17=SSN
    data["spouse"]["first_name"] = cell(19, 5)
    data["spouse"]["middle_name"] = cell(19, 9)
    data["spouse"]["last_name"] = cell(19, 13)
    data["spouse"]["ssn"] = clean_ssn(cell(19, 17))
    # Row 21: Col 14=Spouse Occupation, Col 19=Spouse DOB
    data["spouse"]["occupation"] = cell(21, 14)
    data["spouse"]["dob"] = cell(21, 19)
    # Row 23: Col 3=Spouse Main, Col 9=Spouse Other, Col 15=Spouse Email
    data["spouse"]["main_phone"] = clean_phone(cell(23, 3))
    data["spouse"]["other_phone"] = clean_phone(cell(23, 9))
    data["spouse"]["email"] = cell(23, 15)

    # 3. Mailing Address
    # Row 20: Col 4=Street, Col 11=City, Col 17=State, Col 19=Zip
    data["address"]["street"] = cell(20, 4)
    data["address"]["city"] = cell(20, 11)
    data["address"]["state"] = cell(20, 17)
    data["address"]["zip"] = cell(20, 19)

    # 4. Residential Status & Presence
    # Row 25: Col 5=Residential Status
    data["taxpayer"]["residential_status"] = cell(25, 5)
    data["residency_questions"]["days_2025"] = cell(28, 5)
    data["residency_questions"]["days_2024"] = cell(29, 5)
    data["residency_questions"]["days_2023"] = cell(30, 5)

    # 5. Direct Deposit / Banking Details
    # Col 16 for Rows 27-30
    data["banking"]["bank_name"] = cell(27, 16)
    data["banking"]["routing"] = cell(28, 16)
    data["banking"]["account_number"] = cell(29, 16)
    data["banking"]["account_type"] = cell(30, 16)

    # 6. Dependents (Rows 33 to 38)
    for r in range(33, 39):
        name = cell(r, 2)
        if name:
            dob = cell(r, 7)
            ssn = clean_ssn(cell(r, 11))
            rel = cell(r, 16)
            mos = cell(r, 19) or "12"
            
            # Split first and last name if needed
            name_parts = name.split()
            first_n = name_parts[0] if len(name_parts) > 0 else ""
            last_n = name_parts[-1] if len(name_parts) > 1 else data["taxpayer"].get("last_name", "")

            data["dependents"].append({
                "full_name": name,
                "first_name": first_n,
                "last_name": last_n,
                "dob": dob,
                "ssn": ssn,
                "relationship": rel,
                "months_in_home": mos
            })

    # 7. Residency, FBAR, Digital Asset Questions
    data["residency_questions"]["foreign_account"] = cell(41, 8)
    data["residency_questions"]["foreign_country"] = cell(42, 8)
    data["residency_questions"]["fbar_required"] = cell(41, 39)
    data["residency_questions"]["digital_asset"] = cell(44, 8)
    data["residency_questions"]["health_insurance"] = cell(45, 8)

    # 8. Income Details
    data["income"]["taxpayer_wages"] = cell(64, 8)
    data["income"]["spouse_wages"] = cell(65, 8)
    data["income"]["state_tax_refund"] = cell(63, 18)
    data["income"]["social_security"] = cell(65, 18)
    data["income"]["interest_dividend"] = cell(66, 9)
    data["income"]["unemployment"] = cell(68, 18)
    data["income"]["capital_gains"] = cell(70, 8)
    data["income"]["pension_ira"] = cell(74, 8)

    # 9. Credits & Adjustments
    data["adjustments"]["educator_supplies"] = cell(81, 18)
    data["adjustments"]["tuition_fees"] = cell(82, 8)
    data["adjustments"]["hsa_contributions"] = cell(82, 18)
    data["adjustments"]["child_care_expenses"] = cell(87, 8)
    data["adjustments"]["ira_contributions"] = cell(85, 18)
    data["adjustments"]["student_loan_interest"] = cell(88, 8)

    # 10. Itemized Deductions
    data["deductions"]["prescription_drugs"] = cell(95, 8)
    data["deductions"]["doctors_hospitals"] = cell(100, 8)
    data["deductions"]["real_estate_taxes"] = cell(95, 18)
    data["deductions"]["personal_property_tax"] = cell(96, 18)
    data["deductions"]["home_mortgage_interest"] = cell(103, 8)
    data["deductions"]["charity_contributions"] = cell(103, 18)

    return data

if __name__ == "__main__":
    import pprint
    sample_path = os.path.join(os.path.dirname(__file__), "..", "2025 Personal Tax Organizer.xlsx")
    parsed = parse_organizer(sample_path)
    print("=== SUCCESSFULLY PARSED TAX ORGANIZER ===")
    pprint.pprint(parsed)

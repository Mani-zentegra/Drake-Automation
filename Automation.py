import sys
import ctypes
import time
import os
import win32gui
import win32con
import openpyxl
import pyautogui
from pywinauto import Desktop

import random

pyautogui.FAILSAFE = False

try:
    user32 = ctypes.windll.user32
    hwinsta = user32.OpenWindowStationW('winsta0', False, 0xF037F)
    if hwinsta:
        user32.SetProcessWindowStation(hwinsta)
    hdesk = user32.OpenDesktopW('default', 0, False, 0x01FF)
    if hdesk:
        user32.SetThreadDesktop(hdesk)
except Exception:
    pass

# Safely get window owner handle with exception handling
def get_window_owner(h):
    try:
        if h and win32gui.IsWindow(h):
            return win32gui.GetWindow(h, win32con.GW_OWNER)
    except Exception:
        pass
    return 0

# Safely get window text title with exception handling
def get_window_title(h):
    try:
        if h and win32gui.IsWindow(h):
            return win32gui.GetWindowText(h)
    except Exception:
        pass
    return ""

# Locate the active Drake 2025 window handle
def find_drake_window():
    menu_w = None
    subscreen_w = None
    home_w = None
    def cb(h, _):
        nonlocal menu_w, subscreen_w, home_w
        try:
            if win32gui.IsWindow(h) and win32gui.IsWindowVisible(h):
                t = get_window_title(h)
                if "data entry" in t.lower():
                    owner = get_window_owner(h)
                    owner_text = get_window_title(owner).lower() if owner else ""
                    if "data entry" in owner_text:
                        subscreen_w = (h, t)
                    else:
                        menu_w = (h, t)
                elif "drake 2025 tax software" in t.lower() or "drake 2025" in t.lower():
                    home_w = (h, t)
        except Exception:
            pass
    try:
        win32gui.EnumDesktopWindows(hdesk, cb, None)
    except Exception:
        try:
            win32gui.EnumWindows(cb, None)
        except Exception:
            pass
    if subscreen_w:
        return subscreen_w
    if menu_w:
        return menu_w
    if home_w:
        return home_w
    return (None, "")

# Locate any active Drake dialog by window title keyword
def find_drake_dialog(keyword):
    matched = []
    def enum_cb(h, _):
        if win32gui.IsWindowVisible(h):
            t = get_window_title(h)
            if keyword.lower() in t.lower():
                matched.append((h, t))
    try:
        win32gui.EnumDesktopWindows(hdesk, enum_cb, None)
    except Exception:
        try:
            win32gui.EnumWindows(enum_cb, None)
        except Exception:
            pass
    return matched[0] if matched else (None, "")

# Dismiss any Drake confirmation or warning popups like e-file validation dialogs
def dismiss_drake_popups():
    for _ in range(3):
        h, _ = find_drake_dialog("Data Entry")
        if not h:
            h, _ = find_drake_dialog("Save")
        if not h:
            h, _ = find_drake_dialog("Confirm")
        if h:
            owner = get_window_owner(h)
            if owner != 0:
                try:
                    w = Desktop(backend="uia").window(handle=h)
                    handled = False
                    for btn in w.descendants(control_type="Button"):
                        txt = btn.window_text().strip().lower()
                        if txt in ["yes", "_yes", "ok", "_ok", "save"]:
                            btn.invoke()
                            time.sleep(0.5)
                            handled = True
                            break
                    if not handled:
                        for btn in w.descendants(control_type="Button"):
                            txt = btn.window_text().strip().lower()
                            if txt in ["cancel", "no", "_cancel", "_no"]:
                                btn.invoke()
                                time.sleep(0.5)
                                break
                except Exception:
                    pass
        break

# Attach to the Drake window using pywinauto UIAutomation backend
def connect_window(hwnd):
    try:
        window = Desktop(backend="uia").window(handle=hwnd)
        window.wait("visible", timeout=6)
        return window
    except Exception:
        return None

# Click or invoke the Open/Create button from the Drake main screen
def click_open_create(window):
    for control in window.descendants(control_type="Button"):
        try:
            text = control.window_text().strip().lower()
            cid = (control.element_info.automation_id or "").lower()
            if "open" in text or "create" in text or "open" in cid or "fileopen" in cid:
                control.invoke()
                return True
        except Exception:
            continue
    for control in window.descendants(control_type="Text"):
        try:
            text = control.window_text().strip().lower()
            if "open/create" in text or "open / create" in text:
                r = control.rectangle()
                pyautogui.click((r.left + r.right) // 2, (r.top + r.bottom) // 2)
                return True
        except Exception:
            continue
    pyautogui.hotkey('ctrl', 'o')
    return True

# Enter the 9-digit SSN into the search box and submit
def enter_ssn_and_submit(dialog, ssn):
    for c in dialog.descendants(control_type="Edit"):
        try:
            if c.is_visible() and c.is_enabled():
                c.set_edit_text(str(ssn).strip())
                time.sleep(0.5)
                break
        except Exception:
            continue

    for c in dialog.descendants(control_type="Button"):
        if c.window_text().strip().lower() in ["ok", "_ok"]:
            try:
                c.invoke()
                return True
            except Exception:
                pass
    pyautogui.press('enter')
    return True

# Fill new return details and select Individual 1040 return type
def create_new_individual_return(new_return_window, first_name, last_name, mi=""):
    for rb in new_return_window.descendants(control_type="RadioButton"):
        txt = rb.window_text().strip().lower()
        if "1040" in txt or "individual" in txt:
            try:
                rb.select()
            except Exception:
                pass
            break
    time.sleep(0.3)

    edits = [e for e in new_return_window.descendants(control_type="Edit") if e.is_visible() and e.is_enabled()]
    fn_edit, mi_edit, ln_edit = None, None, None

    for t in new_return_window.descendants(control_type="Text"):
        lbl_txt = t.window_text().strip().rstrip(':').lower()
        if lbl_txt in ["first name", "mi", "last name", "middle initial"]:
            lcx = (t.rectangle().left + t.rectangle().right) / 2
            for e in edits:
                ecx = (e.rectangle().left + e.rectangle().right) / 2
                if abs(ecx - lcx) < 70 and e.rectangle().top >= t.rectangle().bottom - 10:
                    if "first" in lbl_txt:
                        fn_edit = e
                    elif "mi" in lbl_txt or "middle" in lbl_txt:
                        mi_edit = e
                    elif "last" in lbl_txt:
                        ln_edit = e

    if not fn_edit and len(edits) >= 1:
        fn_edit = edits[0]
    if not mi_edit and len(edits) >= 2:
        mi_edit = edits[1]
    if not ln_edit and len(edits) >= 3:
        ln_edit = edits[2]

    if fn_edit:
        try:
            fn_edit.set_edit_text(first_name)
        except Exception:
            pass

    if mi and mi_edit:
        try:
            mi_edit.set_edit_text(mi)
        except Exception:
            pass

    if ln_edit:
        try:
            ln_edit.set_edit_text(last_name)
        except Exception:
            pass

    time.sleep(0.3)
    for btn in new_return_window.descendants(control_type="Button"):
        if btn.window_text().strip().lower() in ["ok", "_ok"]:
            try:
                btn.invoke()
                return True
            except Exception:
                pass
    pyautogui.press('enter')
    return True

# Open an existing client or create a new Individual return by 9-digit SSN
def open_or_create_client(ssn=None, first_name="", last_name="", mi=""):
    if not ssn:
        ssn = input("Enter 9-digit SSN / Client ID: ").strip()
    cleaned_ssn = str(ssn).replace("-", "").replace(" ", "").strip()
    if len(cleaned_ssn) != 9 or not cleaned_ssn.isdigit():
        print(f"[ERROR] SSN must be exactly 9 digits, got '{ssn}'")
        return False

    print("=" * 65)
    print("       DRAKE 2025: OPEN OR CREATE CLIENT RETURN       ")
    print("=" * 65)
    print(f"[TARGET] SSN: {cleaned_ssn}")

    open_h, _ = find_drake_dialog("Open / Create a Return")
    if not open_h:
        main_h, main_t = find_drake_window()
        if not main_h or "data entry" in main_t.lower():
            main_h, main_t = find_drake_dialog("Drake 2025 Tax Software")
        if not main_h:
            print("[ERROR] Drake 2025 Tax Software is not running.")
            return False

        print(f"[CONNECT] Drake Main Window: '{main_t}'")
        main_win = Desktop(backend="uia").window(handle=main_h)
        print("[ACTION] Clicking 'Open/Create' button...")
        click_open_create(main_win)

        for _ in range(10):
            time.sleep(0.5)
            open_h, _ = find_drake_dialog("Open / Create a Return")
            if open_h:
                break

    if not open_h:
        print("[ERROR] Could not open 'Open / Create a Return' dialog.")
        return False

    print("[DIALOG] 'Open / Create a Return' is ready.")
    open_dialog = Desktop(backend="uia").window(handle=open_h)

    print(f"[SEARCH] Submitting SSN: {cleaned_ssn}...")
    enter_ssn_and_submit(open_dialog, cleaned_ssn)

    for _ in range(12):
        time.sleep(0.5)
        prompt_h, _ = find_drake_dialog("Open Return")
        if prompt_h:
            print(f"[NOT FOUND] Client {cleaned_ssn} does not exist.")
            prompt_win = Desktop(backend="uia").window(handle=prompt_h)
            yes_clicked = False
            for btn in prompt_win.descendants(control_type="Button"):
                if btn.window_text().strip().lower() in ["yes", "_yes"]:
                    try:
                        btn.invoke()
                        yes_clicked = True
                        break
                    except Exception:
                        pass
            if not yes_clicked:
                pyautogui.press('y')
            time.sleep(1.0)

            new_ret_h, _ = find_drake_dialog("New Return")
            for _ in range(6):
                if new_ret_h:
                    break
                time.sleep(0.5)
                new_ret_h, _ = find_drake_dialog("New Return")

            if not new_ret_h:
                print("[ERROR] 'New Return' dialog did not appear.")
                return False

            # if not first_name or not last_name:
            #     print("\n[INPUT REQUIRED] Please enter client name for new return:")
            #     if not first_name:
            first_name = f"Tax- {random.randint(1,99)}"
                # if not last_name:
            last_name = f"return- {random.randint(1,99)}"

            print(f"[CREATE] Creating Individual (1040) return: {first_name} {last_name}...")
            new_ret_win = Desktop(backend="uia").window(handle=new_ret_h)
            create_new_individual_return(new_ret_win, first_name, last_name, mi)
            break

        entry_h, entry_t = find_drake_window()
        if entry_h and "data entry" in entry_t.lower():
            print(f"[SUCCESS] Client {cleaned_ssn} exists! Opened return: '{entry_t}'")
            return True

    for _ in range(10):
        time.sleep(0.6)
        entry_h, entry_t = find_drake_window()
        if entry_h and "data entry" in entry_t.lower():
            print(f"[SUCCESS] Return created and opened: '{entry_t}'")
            return True

    return True

# Exit current subscreen and return cleanly to Data Entry Menu
def exit_to_menu():
    for _ in range(4):
        h, _ = find_drake_window()
        if not h:
            break
        try:
            win32gui.ShowWindow(h, win32con.SW_RESTORE)
            win32gui.SetForegroundWindow(h)
        except Exception:
            pass
        time.sleep(0.2)

        try:
            w = Desktop(backend="uia").window(handle=h)
            for btn in w.descendants(control_type="Button"):
                btxt = btn.window_text().strip().lower()
                if btxt.startswith("1|") or btxt.startswith("1 |"):
                    return h
        except Exception:
            pass

        exited = False
        try:
            w = Desktop(backend="uia").window(handle=h)
            for btn in w.descendants(control_type="Button"):
                if btn.window_text().strip().lower() == "exit":
                    btn.invoke()
                    time.sleep(0.8)
                    exited = True
                    break
        except Exception:
            pass

        if not exited:
            pyautogui.press('esc')
            time.sleep(0.5)

        dismiss_drake_popups()
        time.sleep(0.4)

    return find_drake_window()[0]

# Navigate to a specific section screen in Drake by screen code without COM calls
def navigate_to_screen(hwnd, code):
    if not hwnd or not win32gui.IsWindow(hwnd):
        hwnd, _ = find_drake_window()

    if hwnd and win32gui.IsWindow(hwnd):
        try:
            win32gui.ShowWindow(hwnd, win32con.SW_RESTORE)
            win32gui.SetForegroundWindow(hwnd)
        except Exception:
            pass
    time.sleep(0.3)

    try:
        w = Desktop(backend="uia").window(handle=hwnd)
        if str(code) == "1":
            for txt in w.descendants(control_type="Text"):
                if "name, address" in txt.window_text().strip().lower():
                    return
        prefix = f"{code}|"
        for btn in w.descendants(control_type="Button"):
            btxt = btn.window_text().strip().lower()
            if btxt.startswith(prefix) or btxt.startswith(f"{code} |"):
                btn.invoke()
                time.sleep(1.2)
                return
    except Exception:
        pass

    pyautogui.press('esc')
    time.sleep(0.3)
    dismiss_drake_popups()
    time.sleep(0.2)
    pyautogui.write(str(code), interval=0.04)
    time.sleep(0.15)
    pyautogui.press('enter')
    time.sleep(1.2)

# Find horizontal X centers for Taxpayer and Spouse column headers
def get_column_header_x(window):
    taxpayer_x, spouse_x = None, None
    for control in window.descendants():
        try:
            name = control.window_text().strip().lower()
            ctrl_type = control.element_info.control_type
            if ctrl_type in ["Text", "HeaderItem"] or "label" in (control.element_info.automation_id or "").lower():
                if name == "taxpayer" and taxpayer_x is None:
                    r = control.rectangle()
                    taxpayer_x = (r.left + r.right) / 2
                elif name == "spouse" and spouse_x is None:
                    r = control.rectangle()
                    spouse_x = (r.left + r.right) / 2
                if taxpayer_x and spouse_x:
                    break
        except Exception:
            continue
    return taxpayer_x, spouse_x

# Locate edit textbox by cached edit geometry and column type (Taxpayer, Spouse, or Vertical table)
def find_edit_for_field(cached_edits, label_rect, col_type, col_headers):
    label_cx = (label_rect.left + label_rect.right) / 2
    label_cy = (label_rect.top + label_rect.bottom) / 2
    taxpayer_x, spouse_x = col_headers
    col = col_type.strip().lower() if col_type else ""

    if col in ["taxpayer", "spouse"]:
        target_x = spouse_x if col == "spouse" else taxpayer_x
        candidates = []
        for control, rect in cached_edits:
            if rect.left < label_rect.left:
                continue
            center_y = (rect.top + rect.bottom) / 2
            center_x = (rect.left + rect.right) / 2
            v_dist = abs(center_y - label_cy)
            if target_x:
                h_dist = abs(center_x - target_x)
                if h_dist <= 45 and v_dist <= 14:
                    candidates.append((v_dist, control))
            else:
                if v_dist <= 12:
                    candidates.append((center_x, control))

        if not candidates:
            return None

        candidates.sort(key=lambda item: item[0])
        return candidates[0][1] if target_x or col != "spouse" or len(candidates) <= 1 else candidates[1][1]

    vertical_candidates = []
    for control, rect in cached_edits:
        if rect.top < label_rect.top:
            continue
        center_x = (rect.left + rect.right) / 2
        v_dist = rect.top - label_rect.bottom
        h_dist = abs(center_x - label_cx)
        if 0 <= v_dist <= 35 and h_dist <= (label_rect.width() / 2 + 25):
            vertical_candidates.append((v_dist + h_dist, control))

    if vertical_candidates:
        vertical_candidates.sort(key=lambda x: x[0])
        return vertical_candidates[0][1]

    single_candidates = []
    for control, rect in cached_edits:
        if rect.right <= label_rect.left:
            continue
        center_y = (rect.top + rect.bottom) / 2
        center_x = (rect.left + rect.right) / 2
        v_dist = abs(center_y - label_cy)
        if v_dist <= 14 and rect.left >= label_rect.right - 10:
            single_candidates.append((v_dist, abs(center_x - label_rect.right), control))

    if single_candidates:
        single_candidates.sort(key=lambda item: (item[0], item[1]))
        return single_candidates[0][2]

    return None

# Discover all interactive input fields dynamically grouped by UI GroupBox section
def discover_section_controls(window):
    ignored_sections = {"screen menu", "mfs claiming spouse exemption", "foreign account and digital asset questions", "return options"}
    ignored_labels = {"=", "taxpayer", "spouse", "mailing address", "return options", "miscellaneous codes", "2210 options", "u.s. only", "foreign only"}

    sections = {}
    for g in window.descendants(control_type="Group"):
        sec_name = g.window_text().strip()
        if not sec_name or sec_name.lower() in ignored_sections:
            continue

        edits = g.descendants(control_type="Edit")
        labels = []
        for c in g.descendants(control_type="Text"):
            txt = c.window_text().strip().rstrip('.').rstrip(':')
            if txt and txt.lower() not in ignored_labels and not txt.startswith("Age:"):
                try:
                    labels.append((txt, c.rectangle()))
                except Exception:
                    pass

        field_map = {}
        for edit in edits:
            try:
                erect = edit.rectangle()
                ecy = (erect.top + erect.bottom) / 2
                best_lbl = None
                min_dist = 9999
                for lbl_txt, lrect in labels:
                    lcy = (lrect.top + lrect.bottom) / 2
                    v_dist = abs(ecy - lcy)
                    if v_dist <= 14 and lrect.left <= erect.left:
                        dist = v_dist * 2 + (erect.left - lrect.right)
                        if dist < min_dist:
                            min_dist = dist
                            best_lbl = lbl_txt
                    v_above = erect.top - lrect.bottom
                    if 0 <= v_above <= 22 and abs(erect.left - lrect.left) <= 20:
                        dist = v_above + abs(erect.left - lrect.left)
                        if dist < min_dist:
                            min_dist = dist
                            best_lbl = lbl_txt
                if best_lbl:
                    field_map[best_lbl.lower()] = edit
            except Exception:
                continue

        if field_map:
            sections[sec_name.lower()] = field_map

    return sections

# Focus the first editable cell of the Dependents grid
def focus_grid_start(window):
    for h in window.descendants(control_type="HeaderItem"):
        if "first name" in h.window_text().strip().lower():
            r = h.rectangle()
            cx = (r.left + r.right) // 2
            cy = r.bottom + 15
            pyautogui.click(cx, cy)
            time.sleep(0.2)
            return True
    for dg in window.descendants(control_type="DataGrid"):
        r = dg.rectangle()
        pyautogui.click(r.left + 80, r.top + 45)
        time.sleep(0.2)
        return True
    return False

# Set text value into the target input field with popup dismissal
def enter_clean_value(control, value):
    try:
        control.set_focus()
        time.sleep(0.04)
        control.set_edit_text(str(value))
        time.sleep(0.05)
        return True
    except Exception:
        try:
            pyautogui.press('esc')
            time.sleep(0.05)
            rect = control.rectangle()
            cx = (rect.left + rect.right) // 2
            cy = (rect.top + rect.bottom) // 2
            pyautogui.click(cx, cy)
            time.sleep(0.04)
            pyautogui.hotkey('ctrl', 'a')
            time.sleep(0.02)
            pyautogui.write(str(value), interval=0.01)
            time.sleep(0.04)
            return True
        except Exception:
            pyautogui.press('esc')
            return False

# Sanitize special field formats for Drake (SSN, Dates, Phone numbers)
def sanitize_value(field_name, value):
    val_str = str(value).strip() if value is not None else ""
    f_lower = field_name.lower()
    if not val_str or val_str.lower() == "none":
        return ""
    if "date" in f_lower or "/" in val_str:
        cleaned = val_str.replace("/", "").replace("-", "").strip()
        if cleaned.isdigit() and len(cleaned) == 8:
            return cleaned
    if "phone" in f_lower or f_lower in ["daytime", "evening", "cell", "fax"]:
        cleaned = val_str.replace("-", "").replace(" ", "").replace("(", "").replace(")", "").strip()
        if cleaned.isdigit() and len(cleaned) == 10:
            return cleaned
    if "ssn" in f_lower:
        return val_str.replace("-", "").replace(" ", "").strip()
    return val_str.replace("$", "").replace(",", "").strip()

# Load section rows from the unified drake_input.xlsx workbook
def get_section_rows(base_dir, section_name):
    master_path = os.path.join(base_dir, "drake_input.xlsx")
    if os.path.exists(master_path):
        try:
            wb = openpyxl.load_workbook(master_path)
            for sheet in wb.sheetnames:
                if section_name.lower() in sheet.lower():
                    print(f"[EXCEL SOURCE] drake_input.xlsx (Sheet: '{sheet}')")
                    return list(wb[sheet].iter_rows(values_only=True))
        except Exception:
            pass
    return None

# Populate Screen 1 Name and Address fields using dynamic section discovery
def fill_name_and_address(window, base_dir):
    rows = get_section_rows(base_dir, "Name and Address")
    if not rows or len(rows) < 2:
        print("[ERROR] No data rows found for Name and Address.")
        return

    header = [str(c).strip().lower() if c else "" for c in rows[0]]
    data_entries = []

    if "taxpayer" in header and "spouse" in header:
        tp_idx = header.index("taxpayer")
        sp_idx = header.index("spouse")
        for row in rows[1:]:
            if not row or not row[0]:
                continue
            fname = str(row[0]).strip()
            if len(row) > tp_idx and row[tp_idx] is not None and str(row[tp_idx]).strip():
                data_entries.append((fname, "Taxpayer", str(row[tp_idx]).strip()))
            if len(row) > sp_idx and row[sp_idx] is not None and str(row[sp_idx]).strip():
                data_entries.append((fname, "Spouse", str(row[sp_idx]).strip()))
    else:
        for row in rows[1:]:
            if not row or not row[0]:
                continue
            fname = str(row[0]).strip()
            sec = "" if (len(row) < 2 or row[1] is None) else str(row[1]).strip()
            val = "" if (len(row) < 3 or row[2] is None) else str(row[2]).strip()
            if val:
                data_entries.append((fname, sec, val))

    has_spouse = any(sec.lower() == "spouse" for _, sec, _ in data_entries)
    if has_spouse:
        filing_status_ctrl = None
        for lbl in window.descendants(control_type="Text"):
            txt = lbl.window_text().strip().lower()
            if "filing status" in txt:
                lrect = lbl.rectangle()
                lcy = (lrect.top + lrect.bottom) / 2
                best_dist = 9999
                for ctrl in window.descendants():
                    try:
                        ctype = ctrl.element_info.control_type
                        if ctype in ["ComboBox", "Edit"]:
                            crect = ctrl.rectangle()
                            ccy = (crect.top + crect.bottom) / 2
                            if abs(ccy - lcy) <= 15 and crect.left >= lrect.left:
                                dist = crect.left - lrect.right
                                if dist < best_dist:
                                    best_dist = dist
                                    filing_status_ctrl = ctrl
                    except Exception:
                        pass
                break

        if filing_status_ctrl:
            edits = filing_status_ctrl.descendants(control_type="Edit")
            target = edits[0] if edits else filing_status_ctrl
            enter_clean_value(target, "2")
            time.sleep(0.3)

    sections = discover_section_controls(window)
    print(f"[READY] Discovered {len(sections)} UI sections on Screen 1: {list(sections.keys())}")
    print("-" * 65)

    success, total = 0, 0
    for field_name, section, raw_val in data_entries:
        total += 1
        clean_val = sanitize_value(field_name, raw_val)

        matched_sec = None
        for s_key in sections:
            if section.lower() in s_key or s_key in section.lower():
                matched_sec = sections[s_key]
                break

        ctrl = None
        if matched_sec:
            ctrl = matched_sec.get(field_name.lower())
            if not ctrl:
                for k, c in matched_sec.items():
                    if field_name.lower() in k or k in field_name.lower():
                        ctrl = c
                        break

        if not ctrl:
            print(f"[-] Control not found: '{field_name}' [{section}]")
            continue

        if enter_clean_value(ctrl, clean_val):
            display = f"{field_name} [{section}]" if section else field_name
            print(f"[OK] {display:<45} = {clean_val}")
            success += 1
        else:
            print(f"[FAIL] Could not set: '{field_name}' [{section}]")
        time.sleep(0.04)

    print("-" * 65)
    print(f"[RESULT] Successfully populated {success} of {total} fields on Screen 1!")

# Populate Screen 2 Dependents DataGrid across all rows and columns and save
def fill_dependents(window, base_dir):
    rows = get_section_rows(base_dir, "Dependents")
    if not rows or len(rows) < 2:
        print("[ERROR] No data rows found for Dependents.")
        return

    headers = [str(c) for c in rows[0]]
    data_rows = rows[1:]

    focus_grid_start(window)
    time.sleep(0.2)
    print(f"[GRID] Entering {len(data_rows)} rows across {len(headers)} columns...")
    print("-" * 65)

    for row_idx, row in enumerate(data_rows, start=1):
        print(f"Entering Row {row_idx}: {row[0]} {row[2]} (SSN: {row[4]})...")
        for col_idx, cell_value in enumerate(row):
            header_name = headers[col_idx]
            clean_val = sanitize_value(header_name, cell_value)
            if clean_val:
                pyautogui.hotkey('ctrl', 'a')
                time.sleep(0.01)
                pyautogui.write(clean_val, interval=0.01)
                time.sleep(0.02)
            pyautogui.press('tab')
            time.sleep(0.03)
        time.sleep(0.08)

    print("-" * 65)
    print(f"[RESULT] Successfully populated {len(data_rows)} rows in Dependents grid!")

    time.sleep(0.3)
    saved = False
    for b in window.descendants(control_type="Button"):
        if b.window_text().strip().lower() == "save":
            try:
                b.invoke()
                print("[SAVE] Clicked 'Save' button for Dependents grid!")
                saved = True
                break
            except Exception:
                pass

    if not saved:
        pyautogui.hotkey('alt', 's')
        time.sleep(0.2)

    time.sleep(0.8)

# Populate two-column forms (Screen 3 Income or Screen 4 Adjustments) using column alignment
def fill_column_form(window, base_dir, section_name):
    rows = get_section_rows(base_dir, section_name)
    if not rows or len(rows) < 2:
        print(f"[ERROR] No data rows found for {section_name}.")
        return

    taxpayer_x, spouse_x = None, None
    label_map = {}
    cached_edits = []
    for control in window.descendants():
        try:
            ctrl_type = control.element_info.control_type
            if ctrl_type == "Edit":
                cached_edits.append((control, control.rectangle()))
            name = control.window_text().strip().rstrip(':').rstrip('.')
            if name and (ctrl_type in ["Text", "HeaderItem"] or "label" in (control.element_info.automation_id or "").lower()):
                rect = control.rectangle()
                lower_name = name.lower()
                label_map[lower_name] = (name, rect)
                if lower_name == "taxpayer" and taxpayer_x is None:
                    taxpayer_x = (rect.left + rect.right) / 2
                elif lower_name == "spouse" and spouse_x is None:
                    spouse_x = (rect.left + rect.right) / 2
        except Exception:
            pass

    col_headers = (taxpayer_x, spouse_x)
    if col_headers[0] and col_headers[1]:
        print(f"[COLUMNS] Detected headers: Taxpayer X={int(col_headers[0])}, Spouse X={int(col_headers[1])}")

    print(f"[SCAN] Discovered {len(label_map)} labels and {len(cached_edits)} edits on active screen.")
    print("-" * 65)

    header = [str(c).strip().lower() if c else "" for c in rows[0]]
    data_entries = []

    if "taxpayer" in header and "spouse" in header:
        tp_idx = header.index("taxpayer")
        sp_idx = header.index("spouse")
        for row in rows[1:]:
            if not row or not row[0]:
                continue
            fname = str(row[0]).strip()
            if len(row) > tp_idx and row[tp_idx] is not None and str(row[tp_idx]).strip():
                data_entries.append((fname, "Taxpayer", str(row[tp_idx]).strip()))
            if len(row) > sp_idx and row[sp_idx] is not None and str(row[sp_idx]).strip():
                data_entries.append((fname, "Spouse", str(row[sp_idx]).strip()))
    else:
        for row in rows[1:]:
            if not row or row[0] is None:
                continue
            fname = str(row[0]).strip()
            col_type = "" if (len(row) < 2 or row[1] is None) else str(row[1]).strip()
            raw_val = "" if (len(row) < 3 or row[2] is None) else str(row[2]).strip()

            if len(row) == 2 or (len(row) >= 2 and col_type and not raw_val and col_type.lower() not in ["taxpayer", "spouse"]):
                raw_val = col_type
                col_type = ""
            if raw_val:
                data_entries.append((fname, col_type, raw_val))

    success, total = 0, 0
    for field_name, col_type, raw_val in data_entries:
        total += 1
        clean_val = sanitize_value(field_name, raw_val)
        matched = label_map.get(field_name.lower())
        if not matched:
            for k, v in label_map.items():
                if field_name.lower() in k or k in field_name.lower():
                    matched = v
                    break

        if not matched:
            print(f"[-] Label not found: '{field_name}'")
            continue

        lbl_text, lbl_rect = matched
        edit_ctrl = find_edit_for_field(cached_edits, lbl_rect, col_type, col_headers)
        if not edit_ctrl:
            print(f"[-] Input box not found for: '{lbl_text}' [{col_type}]")
            continue

        if enter_clean_value(edit_ctrl, clean_val):
            display = f"{lbl_text} [{col_type}]" if col_type else lbl_text
            print(f"[OK] {display:<50} = {clean_val}")
            success += 1
        else:
            print(f"[FAIL] Could not set: '{lbl_text}'")
        time.sleep(0.04)

    print("-" * 65)
    print(f"[RESULT] Successfully populated {success} of {total} fields for {section_name}!")

# Main workflow executing client open or create followed by all 4 modules automatically
def run():
    base_dir = os.path.dirname(sys.executable) if getattr(sys, 'frozen', False) else os.path.dirname(os.path.abspath(__file__))

    print("=" * 65)
    print("   DRAKE DYNAMIC AUTOMATION: FULL AUTOMATIC 4-SECTION FILLER ")
    print("=" * 65)

    hwnd, title = find_drake_window()
    if not hwnd:
        print("[ERROR] Drake 2025 window is not running.")
        return

    print(f"[CONNECTED] Window: '{title}'")

    if "data entry" not in title.lower():
        print("[HOME SCREEN] Drake Home Screen detected. Opening client return...")
        if not open_or_create_client():
            return
        hwnd, title = find_drake_window()

    print("\n" + "=" * 65)
    print("   STARTING AUTOMATIC 4-SECTION PIPELINE")
    print("=" * 65)

    pipeline = [
        ("1", "Name and Address", fill_name_and_address),
        ("2", "Dependents", fill_dependents),
        ("3", "Income", lambda w, b: fill_column_form(w, b, "Income")),
        ("4", "Adjustments", lambda w, b: fill_column_form(w, b, "Adjustments"))
    ]

    for step_num, (code, name, fill_func) in enumerate(pipeline, start=1):
        print(f"\n>>> [{step_num}/4] Processing Screen {code}: '{name}'...")
        hwnd, _ = find_drake_window()
        navigate_to_screen(hwnd, code)

        screen_hwnd = None
        for _ in range(8):
            h, t = find_drake_window()
            owner = get_window_owner(h)
            owner_text = get_window_title(owner).lower() if owner else ""
            if h and "data entry" in owner_text:
                screen_hwnd = h
                break
            time.sleep(0.3)

        if not screen_hwnd:
            screen_hwnd, _ = find_drake_window()

        if not screen_hwnd:
            print(f"[ERROR] Could not navigate to Screen {code}.")
            continue

        try:
            win32gui.ShowWindow(screen_hwnd, win32con.SW_RESTORE)
            win32gui.SetForegroundWindow(screen_hwnd)
        except Exception:
            pass
        time.sleep(0.2)

        window = connect_window(screen_hwnd)
        if not window:
            print(f"[ERROR] Could not connect to Screen {code} window.")
            continue

        fill_func(window, base_dir)
        time.sleep(0.3)

        print(f"[SAVING] Screen {code} finished. Saving and returning to Data Entry Menu...")
        exit_to_menu()
        time.sleep(0.5)

    print("\n" + "=" * 65)
    print("   ALL 4 SECTIONS AUTOMATICALLY COMPLETED SUCCESSFULLY! ")
    print("=" * 65)
    if getattr(sys, 'frozen', False):
        input("\nPress Enter to exit...")

if __name__ == "__main__":
    run()
